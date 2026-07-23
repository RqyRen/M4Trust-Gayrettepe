"""DOCUMENT_EXTRACTION pipeline testleri (ADR-002 §3.1, §11, §13; ADR-001 §16).

Indirme, hash, tur tespiti, GERCEK PDF/DOCX metin cikarimi, canonical mapping
ve teknik schema validation gercek calisir. LLM cagrisi (ADR-004 §17
"mock-first") monkeypatch ile sahtelenir; gercek GPT-5.4 entegrasyonu ayri bir
canli script ile (pytest disinda) dogrulanir.
"""
from __future__ import annotations

import hashlib
import io
import json
import threading
from datetime import datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, HTTPServer
from pathlib import Path

import docx
import fitz  # PyMuPDF
import pytest
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfWriter

from app.config import Settings
from app.contracts.errors import ContractViolation, ErrorCode, PipelineFailure
from app.contracts.validation import validate_outgoing
from app.pipeline.document_extraction import llm as llm_module
from app.pipeline.document_extraction import mapping as mapping_module
from app.pipeline.document_extraction import name_masking as name_masking_module
from app.pipeline.document_extraction.media import DOCX, PDF, detect_media_type
from app.pipeline.document_extraction import pipeline as pipeline_module
from app.pipeline.document_extraction.pipeline import run
from app.pipeline.registry import pipeline_for

_EXAMPLES = Path(__file__).resolve().parents[1] / "contracts" / "examples"


def _real_pdf_bytes(text: str = "ACME Corp agrees to pay 1000 EUR within 30 days.") -> bytes:
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(text=text)
    return bytes(pdf.output())


def _real_docx_bytes(text: str = "Delivery of sealed boxes is required with a signed delivery note.") -> bytes:
    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _encrypted_pdf_bytes() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt(user_password="secret")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _text_image_bytes(lines: list[str], *, size: tuple[int, int] = (1200, 300)) -> bytes:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 40)
    except OSError:
        font = ImageFont.load_default()
    for i, line in enumerate(lines):
        draw.text((40, 40 + i * 70), line, fill="black", font=font)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _scanned_pdf_bytes(lines: list[str]) -> bytes:
    """Metin katmani OLMAYAN, sadece goruntu iceren bir sayfa - gercek
    taranmis/fotograflanmis belgeyi simule eder."""
    width, height = 1200, 300
    image_bytes = _text_image_bytes(lines, size=(width, height))
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    page.insert_image(fitz.Rect(0, 0, width, height), stream=image_bytes)
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _blank_scanned_pdf_bytes() -> bytes:
    """Tamamen bos (metinsiz) taranmis sayfa - OCR sonrasi da bos kalmali."""
    width, height = 400, 300
    image = Image.new("RGB", (width, height), "white")
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=width, height=height)
    page.insert_image(fitz.Rect(0, 0, width, height), stream=buffer.getvalue())
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


def _hybrid_pdf_bytes() -> bytes:
    """1. sayfa dijital metin, 2. sayfa taranmis goruntu - karisik belge."""
    digital = PdfWriter(clone_from=io.BytesIO(_real_pdf_bytes("Digital cover page with real text layer.")))
    scanned = PdfWriter(clone_from=io.BytesIO(_scanned_pdf_bytes(["SCANNED PAGE: DELIVERY TERMS"])))
    digital.append(scanned)
    buffer = io.BytesIO()
    digital.write(buffer)
    return buffer.getvalue()


_PDF_BYTES = _real_pdf_bytes()
_DOCX_BYTES = _real_docx_bytes()
_ENCRYPTED_PDF_BYTES = _encrypted_pdf_bytes()
_SCANNED_PDF_BYTES = _scanned_pdf_bytes(["PAYMENT: Buyer pays 1000 EUR", "within 30 days of invoice date."])
_BLANK_SCANNED_PDF_BYTES = _blank_scanned_pdf_bytes()
_HYBRID_PDF_BYTES = _hybrid_pdf_bytes()
_JUNK = b"just plain text, not a document"
_MALFORMED_PDF = b"%PDF-1.7\n" + b"not a real pdf body, no xref table " * 20
_PDF_WITH_TAX_ID_BYTES = _real_pdf_bytes("ACME Corp, Vergi No: 1234567890, agrees to pay 1000 EUR within 30 days.")
_PDF_WITH_PERSON_NAME_BYTES = _real_pdf_bytes("This contract is signed by Ahmet Yilmaz.")

_BODIES = {
    "/pdf": _PDF_BYTES,
    "/docx": _DOCX_BYTES,
    "/junk": _JUNK,
    "/malformed-pdf": _MALFORMED_PDF,
    "/encrypted": _ENCRYPTED_PDF_BYTES,
    "/scanned": _SCANNED_PDF_BYTES,
    "/blank-scanned": _BLANK_SCANNED_PDF_BYTES,
    "/hybrid": _HYBRID_PDF_BYTES,
    "/pdf-with-tax-id": _PDF_WITH_TAX_ID_BYTES,
    "/pdf-with-person-name": _PDF_WITH_PERSON_NAME_BYTES,
}


class _Handler(BaseHTTPRequestHandler):
    def do_GET(self) -> None:  # noqa: N802
        body = _BODIES.get(self.path)
        if body is None:
            self.send_response(404)
            self.end_headers()
            return
        self.send_response(200)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *args) -> None:
        return


@pytest.fixture(autouse=True)
def _no_real_legal_rag(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bu dosyadaki testler legal grounding'i degil, extraction'i test ediyor.

    legal_rag.retrieve() mock'lanmazsa gercek BGE-M3 modelini belleye yukler
    (agir, ~saniyeler suren bir islem) -- bu dosyadaki her test bunu bilmeden
    odemis olur. Varsayilan olarak bos sonuc donduruyoruz; retrieval'in
    kendisini test eden test bunu kendi icinde ayrica override eder.
    """
    monkeypatch.setattr(llm_module.legal_rag, "retrieve", lambda *a, **k: [])
    monkeypatch.setattr(llm_module.legal_rag, "retrieve_batch", lambda query_texts, top_k=5: [[] for _ in query_texts])


@pytest.fixture(autouse=True)
def _no_real_name_masking(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bu dosyadaki testler NER tabanli kisi-adi maskelemeyi degil, extraction'i
    test ediyor. Mock'lanmazsa her testte agir bir transformers NER modeli
    belleye yuklenir (bkz. _no_real_legal_rag ile ayni gerekce). Bu ozelligi
    kendisi test eden test (asagida) bu fixture'i kendi icinde override eder.
    """
    monkeypatch.setattr(pipeline_module, "mask_person_names", lambda text: (text, {}))
    monkeypatch.setattr(pipeline_module, "restore_person_names", lambda value, name_map: value)


@pytest.fixture(scope="module")
def base_url() -> str:
    server = HTTPServer(("127.0.0.1", 0), _Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


def _request(base_url: str, path: str, body: bytes, media_type: str = PDF) -> dict:
    req = json.loads((_EXAMPLES / "document-extraction" / "full-request.json").read_text(encoding="utf-8"))
    expires = (datetime.now(timezone.utc) + timedelta(minutes=10)).strftime("%Y-%m-%dT%H:%M:%SZ")
    req["payload"]["input"].update(
        {
            "mediaType": media_type,
            "sizeBytes": len(body),
            "sha256": hashlib.sha256(body).hexdigest(),
            "download": {"url": f"{base_url}{path}", "expiresAt": expires},
        }
    )
    # Fixture'daki sabit deadlineAt zamanla gecmise duser (Berke review #5:
    # deadline artik semantik kontrol ediliyor) -- testler her zaman calisan
    # zamana gore GELECEK bir deadline kullanmali.
    req["payload"]["deadlineAt"] = (datetime.now(timezone.utc) + timedelta(minutes=30)).strftime("%Y-%m-%dT%H:%M:%SZ")
    return req


_CANNED_LLM_OUTPUT = {
    "detectedLanguage": "en",
    "parties": [
        {
            "role": "SELLER",
            "legalName": "ACME Corp",
            "legalNameConfidence": 0.95,
            "taxIdentifier": "1234567890",  # mapping.py bunu daima maskelemeli
            "taxIdentifierConfidence": 0.8,
            "page": 1,
        }
    ],
    "rules": [
        {
            "category": "PAYMENT",
            "title": "Payment amount",
            "description": "Buyer pays 1000 EUR within 30 days.",
            "valueType": "MONEY",
            "textValue": None,
            "amountMinor": 100000,
            "currency": "eur",
            "basisPoints": None,
            "durationDays": None,
            "dateValue": None,
            "booleanValue": None,
            "quantityValue": None,
            "quantityUnit": None,
            "confidence": 0.9,
            "page": 1,
        }
    ],
    "deliveryRequirements": [
        {"evidenceType": "DELIVERY_NOTE", "required": True, "confidence": 0.85, "page": 1}
    ],
    "requiresManualReview": False,
    "reviewReasons": [],
}


def _mock_llm(monkeypatch: pytest.MonkeyPatch, output: dict = _CANNED_LLM_OUTPUT) -> None:
    monkeypatch.setattr(llm_module, "extract_structured_data", lambda text, settings: output)


# --- LLM cagri parametreleri (gercek API'ye gitmeden, OpenAI client mock) ---

def test_llm_call_uses_temperature_zero_for_output_consistency(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bulgu (15 Temmuz 2026): temperature ayarlanmamisti; ayni sozlesme
    calistirma-calistirmaya farkli warning sayisi uretebiliyordu. Bu test
    gercek OpenAI cagrisi yapmadan, `temperature=0`'in fiilen API'ye
    gonderilen kwarg'larda oldugunu dogrular.
    """
    captured_kwargs: dict = {}

    class _FakeMessage:
        content = json.dumps(_CANNED_LLM_OUTPUT)

    class _FakeCompletions:
        def create(self, **kwargs):
            captured_kwargs.update(kwargs)
            return type("_FakeResponse", (), {"choices": [type("_FakeChoice", (), {"message": _FakeMessage()})()]})()

    class _FakeChat:
        completions = _FakeCompletions()

    class _FakeClient:
        def __init__(self, **kwargs):
            pass

        chat = _FakeChat()

    monkeypatch.setattr(llm_module, "OpenAI", _FakeClient)
    settings = Settings(openai_api_key="test-key", openai_model="gpt-5.4")

    llm_module.extract_structured_data("sample contract text", settings)

    assert captured_kwargs["temperature"] == 0


def test_legal_rag_retrieval_failure_does_not_break_extraction(monkeypatch: pytest.MonkeyPatch) -> None:
    """Legal RAG bir kalite artiricidir, extraction'in on kosulu degil: kanun
    kulliyati embed edilmemis olsa veya retrieval baska bir sebeple patlasa
    bile extract_structured_data eskisi gibi basariyla sonuc uretmeye devam
    etmeli, sadece baglamsiz kalmali.
    """
    captured_kwargs: dict = {}

    class _FakeMessage:
        content = json.dumps(_CANNED_LLM_OUTPUT)

    class _FakeCompletions:
        def create(self, **kwargs):
            captured_kwargs.update(kwargs)
            return type("_FakeResponse", (), {"choices": [type("_FakeChoice", (), {"message": _FakeMessage()})()]})()

    class _FakeChat:
        completions = _FakeCompletions()

    class _FakeClient:
        def __init__(self, **kwargs):
            pass

        chat = _FakeChat()

    def _broken_retrieve(*args, **kwargs):
        raise FileNotFoundError("legal corpus embeddings not built")

    monkeypatch.setattr(llm_module, "OpenAI", _FakeClient)
    monkeypatch.setattr(llm_module.legal_rag, "retrieve", _broken_retrieve)
    settings = Settings(openai_api_key="test-key", openai_model="gpt-5.4")

    result = llm_module.extract_structured_data("sample contract text", settings)

    assert result == _CANNED_LLM_OUTPUT
    # Legal context sistem mesaji hic eklenmemis olmali -- sadece SYSTEM_PROMPT + user mesaji.
    assert len(captured_kwargs["messages"]) == 2


def test_map_rule_attaches_legal_basis_for_a_confident_match(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        mapping_module.legal_rag,
        "retrieve",
        lambda query, top_k=1: [{"source": "tbk-6098", "madde_no": "179", "score": 0.71, "text": "..."}],
    )
    rule = mapping_module._map_rule(
        {"category": "PENALTY", "title": "Gecikme cezasi", "description": "Gunluk binde bir gecikme cezasi.", "confidence": 0.9},
        0,
        warnings=[],
    )
    assert rule["legalBasis"] == {"source": "tbk-6098", "articleNo": "179"}


def test_map_rule_omits_legal_basis_below_score_threshold(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(
        mapping_module.legal_rag,
        "retrieve",
        lambda query, top_k=1: [{"source": "tbk-6098", "madde_no": "1", "score": 0.1, "text": "..."}],
    )
    rule = mapping_module._map_rule(
        {"category": "OTHER", "title": "Alakasiz madde", "description": "...", "confidence": 0.5}, 0, warnings=[]
    )
    assert "legalBasis" not in rule


def test_map_rule_omits_legal_basis_for_unknown_source(monkeypatch: pytest.MonkeyPatch) -> None:
    """legalBasis.source semada kapali bir enum -- kulliyat disi bir kaynak asla ciktiya sizmamali."""
    monkeypatch.setattr(
        mapping_module.legal_rag,
        "retrieve",
        lambda query, top_k=1: [{"source": "unexpected-source", "madde_no": "1", "score": 0.9, "text": "..."}],
    )
    rule = mapping_module._map_rule(
        {"category": "OTHER", "title": "Test", "description": "...", "confidence": 0.5}, 0, warnings=[]
    )
    assert "legalBasis" not in rule


def test_map_rule_omits_legal_basis_when_retrieval_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    """legal_rag cokse bile (embedding yok, model yuklenemedi vb.) rule basariyla eslenmeli."""

    def _broken_retrieve(*args, **kwargs):
        raise FileNotFoundError("legal corpus embeddings not built")

    monkeypatch.setattr(mapping_module.legal_rag, "retrieve", _broken_retrieve)
    rule = mapping_module._map_rule(
        {"category": "PAYMENT", "title": "Odeme", "description": "30 gun icinde odeme.", "confidence": 0.8}, 0, warnings=[]
    )
    assert "legalBasis" not in rule
    assert rule["category"] == "PAYMENT"


def test_map_to_canonical_result_batches_legal_basis_lookup_in_a_single_call(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bulgu (23 Temmuz 2026, Railway'de kanitlandi): N kural icin N ayri
    legal_rag.retrieve() cagrisi, OpenAI cevap verdikten SONRA kural sayisiyla
    orantili bir gecikme yaratiyordu (12 kuralli bir sozlesmede ~7 dakika).
    map_to_canonical_result artik TUM kurallari TEK bir retrieve_batch()
    cagrisinda islemeli -- kac kural olursa olsun.
    """
    calls: list[list[str]] = []

    def _fake_retrieve_batch(query_texts, top_k=1):
        calls.append(list(query_texts))
        return [
            [{"source": "tbk-6098", "madde_no": str(i + 1), "score": 0.9, "text": "..."}]
            for i in range(len(query_texts))
        ]

    monkeypatch.setattr(mapping_module.legal_rag, "retrieve_batch", _fake_retrieve_batch)

    llm_output = {
        "parties": [],
        "rules": [
            {"category": "PAYMENT", "title": f"Kural {i}", "description": "...", "confidence": 0.9}
            for i in range(5)
        ],
        "deliveryRequirements": [],
        "requiresManualReview": False,
        "reviewReasons": [],
    }
    result, _ = mapping_module.map_to_canonical_result(llm_output, document={})

    assert len(calls) == 1  # tek batch cagrisi, kural basina degil
    assert len(calls[0]) == 5
    assert len(result["rules"]) == 5
    for i, rule in enumerate(result["rules"]):
        assert rule["legalBasis"] == {"source": "tbk-6098", "articleNo": str(i + 1)}


def test_map_to_canonical_result_omits_legal_basis_when_batch_retrieval_fails(monkeypatch: pytest.MonkeyPatch) -> None:
    def _broken_retrieve_batch(query_texts, top_k=1):
        raise FileNotFoundError("legal corpus embeddings not built")

    monkeypatch.setattr(mapping_module.legal_rag, "retrieve_batch", _broken_retrieve_batch)

    llm_output = {
        "parties": [],
        "rules": [{"category": "PAYMENT", "title": "Odeme", "description": "30 gun icinde odeme.", "confidence": 0.8}],
        "deliveryRequirements": [],
        "requiresManualReview": False,
        "reviewReasons": [],
    }
    result, _ = mapping_module.map_to_canonical_result(llm_output, document={})

    assert len(result["rules"]) == 1
    assert "legalBasis" not in result["rules"][0]


def _fake_openai_client(*, choices: list):
    class _FakeCompletions:
        def create(self, **kwargs):
            return type("_FakeResponse", (), {"choices": choices})()

    class _FakeChat:
        completions = _FakeCompletions()

    class _FakeClient:
        def __init__(self, **kwargs):
            pass

        chat = _FakeChat()

    return _FakeClient


def test_malformed_json_response_becomes_stable_pipeline_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    """Bulgu (17 Temmuz 2026, Berke review #9): bozuk JSON dogrudan json.loads()'a
    veriliyordu, ham JSONDecodeError firlatiyordu -- run_with_retry bunu
    PipelineFailure SANMADIGI icin yakalamiyordu ve Spring'e HICBIR
    ai.job.failed.v1 uretilmiyordu (mesaj sessizce dead-letter'a dusuyordu).
    """
    message = type("_FakeMessage", (), {"content": "{not valid json"})()
    choice = type("_FakeChoice", (), {"message": message})()
    monkeypatch.setattr(llm_module, "OpenAI", _fake_openai_client(choices=[choice]))
    settings = Settings(openai_api_key="test-key", openai_model="gpt-5.4")

    with pytest.raises(PipelineFailure) as exc_info:
        llm_module.extract_structured_data("sample contract text", settings)

    assert exc_info.value.code == ErrorCode.MODEL_PROVIDER_UNAVAILABLE


def test_empty_choices_response_becomes_stable_pipeline_failure(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(llm_module, "OpenAI", _fake_openai_client(choices=[]))
    settings = Settings(openai_api_key="test-key", openai_model="gpt-5.4")

    with pytest.raises(PipelineFailure) as exc_info:
        llm_module.extract_structured_data("sample contract text", settings)

    assert exc_info.value.code == ErrorCode.MODEL_PROVIDER_UNAVAILABLE


# --- Pipeline uctan uca (gercek indirme + hash + tur + GERCEK metin cikarimi, LLM mock) ---

def test_pdf_pipeline_produces_schema_valid_completed_event(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _mock_llm(monkeypatch)
    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)

    assert event["eventType"] == "ai.job.completed.v1"
    assert event["jobType"] == "DOCUMENT_EXTRACTION"
    assert event["causationId"] == request["eventId"]
    validate_outgoing(event)

    document = event["payload"]["result"]["document"]
    assert document["detectedMediaType"] == PDF
    assert document["textExtractionMethod"] == "DIGITAL_PDF"
    assert document["pageCount"] == 1
    assert document["contentSha256"] == hashlib.sha256(_PDF_BYTES).hexdigest()

    # Bulgu (21 Temmuz 2026, bagimsiz denetim): retrievalVersion sabit None
    # yaziyordu, RAG (PR #31) gercekten calisiyor olmasina ragmen. Artik
    # pipelineVersion/promptVersion ile ayni statik-versiyonlama desenini izler.
    assert event["payload"]["technicalMetadata"]["retrievalVersion"] == pipeline_module.LEGAL_RAG_VERSION


def test_llm_never_receives_raw_tax_id_but_output_shape_is_unchanged(
    base_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Bulgu (17 Temmuz 2026): ham metin (vergi no dahil) LLM'e maskesiz gidiyordu.

    Gercek bir PDF'e gercek bir vergi no gomer, gercek indirme+parse'tan sonra
    LLM'e ULASAN metni yakalar (spy), ham numaranin ORADA hic bulunmadigini
    kanitlar -- canonical cikti sekli/semasi ise degismeden ayni kalir.
    """
    tax_id = "1234567890"

    captured_text: dict[str, str] = {}

    def _spy_llm(text: str, settings):
        captured_text["value"] = text
        return _CANNED_LLM_OUTPUT

    monkeypatch.setattr(llm_module, "extract_structured_data", _spy_llm)
    request = _request(base_url, "/pdf-with-tax-id", _PDF_WITH_TAX_ID_BYTES)
    event = run(request)

    assert tax_id not in captured_text["value"]
    assert "[MASKED_TAX_ID]" in captured_text["value"]

    # Canonical cikti sekli hic degismedi: ayni alanlar, hala schema-valid,
    # vergi kimligi burada zaten (mapping.py yuzunden) her zaman maskeliydi.
    validate_outgoing(event)
    assert event["payload"]["result"]["parties"][0]["taxIdentifier"] == {
        "value": None,
        "masked": True,
        "confidence": 0.8,
    }


def test_person_name_masked_before_llm_and_restored_in_party_legal_name(
    base_url: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    """20 Temmuz 2026: name_masking.py'nin pipeline'a dogru bagli oldugunu kanitlar.

    NER modelinin kendisi burada sahtelenir (gercegi test_name_masking.py'de
    ayrica test edildi); bu test sadece round-trip'i kanitlar: (1) token'li
    metin LLM'e gider, gercek isim gitmez, (2) LLM ciktisinda partiy'nin
    legalName'i olarak KULLANILAN token, canonical sonucta gercek isme geri
    donusturulur -- taraf-adi cikarimi bozulmaz.
    """
    real_name = "Ahmet Yilmaz"
    token = "[MASKED_PERSON_1]"

    def _fake_mask(text: str) -> tuple[str, dict[str, str]]:
        assert real_name in text
        return text.replace(real_name, token), {token: real_name}

    monkeypatch.setattr(pipeline_module, "mask_person_names", _fake_mask)
    monkeypatch.setattr(pipeline_module, "restore_person_names", name_masking_module.restore_person_names)

    captured_text: dict[str, str] = {}

    def _spy_llm(text: str, settings):
        captured_text["value"] = text
        return {
            **_CANNED_LLM_OUTPUT,
            "parties": [
                {
                    "role": "SELLER",
                    "legalName": token,
                    "legalNameConfidence": 0.9,
                    "taxIdentifier": None,
                    "taxIdentifierConfidence": 0.0,
                    "page": 1,
                }
            ],
        }

    monkeypatch.setattr(llm_module, "extract_structured_data", _spy_llm)
    request = _request(base_url, "/pdf-with-person-name", _PDF_WITH_PERSON_NAME_BYTES)
    event = run(request)

    assert real_name not in captured_text["value"]
    assert token in captured_text["value"]

    assert event["payload"]["result"]["parties"][0]["legalName"]["value"] == real_name


def test_truncated_document_produces_warning_and_forces_review(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    # ADR-002 §13: prompt sinirini asan belgede sessiz veri kaybi OLMAMALI.
    # Kanit: 15 Temmuz 2026 kalite testinde, kesme yuzunden gercek bir
    # 750.000 TL'lik ceza maddesi warning'siz kayboluyordu (bkz. pipeline.py
    # _truncation_warning docstring'i). Bu test o senaryoyu simule eder.
    from app.config import get_settings

    _mock_llm(monkeypatch)
    settings = get_settings()
    monkeypatch.setattr(settings, "openai_max_source_chars", 50)  # _PDF_BYTES metninden kucuk

    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)

    validate_outgoing(event)
    warnings = event["payload"]["warnings"]
    assert any(w["code"] == "PARTIAL_TEXT_EXTRACTION" for w in warnings)

    summary = event["payload"]["result"]["summary"]
    assert summary["requiresManualReview"] is True
    assert any("processing limit" in reason for reason in summary["reviewReasons"])


def test_check_cancelled_before_llm_call_stops_pipeline(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    """Berke review #1: pahali LLM cagrisindan ONCE cancellation kontrol
    edilmeli. Ilk checkpoint (indirmeden once) gecerli sayilir, ikinci
    checkpoint'te (LLM'den once) iptal simule edilir -- indirme/parse GERCEK
    calisir ama LLM'e hic ulasilmamalidir.
    """
    from app.common.cancellation import JobCancelled

    calls = {"n": 0}

    def _check_cancelled() -> None:
        calls["n"] += 1
        if calls["n"] == 2:
            raise JobCancelled("job-under-test")

    llm_calls: list[int] = []
    monkeypatch.setattr(llm_module, "extract_structured_data", lambda text, settings: llm_calls.append(1))

    request = _request(base_url, "/pdf", _PDF_BYTES)
    with pytest.raises(JobCancelled):
        run(request, check_cancelled=_check_cancelled)

    assert llm_calls == []  # LLM cagrisi hic yapilmadi -- maliyet onlendi
    assert calls["n"] == 2


def test_docx_pipeline_produces_schema_valid_completed_event(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _mock_llm(monkeypatch)
    request = _request(base_url, "/docx", _DOCX_BYTES, media_type=DOCX)
    event = run(request)

    validate_outgoing(event)
    document = event["payload"]["result"]["document"]
    assert document["detectedMediaType"] == DOCX
    assert document["textExtractionMethod"] == "DOCX"


def test_tax_identifier_is_always_masked(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    # ADR-001 §16: ozel bir maskeleme alt sistemi olmadan ham deger asla cikmaz.
    _mock_llm(monkeypatch)
    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)
    party = event["payload"]["result"]["parties"][0]
    assert party["taxIdentifier"]["masked"] is True
    assert party["taxIdentifier"]["value"] is None


def test_duplicate_party_mentions_are_merged(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    # Bulgu (gercek sozlesme testi, 15 Temmuz 2026): ayni sirket baslikta ve
    # imza blogunda hafif farkli bicimlendirmeyle gectiginde LLM iki ayri
    # taraf cikarabiliyordu (ornek: "LTD.STI." vs "LTD. STI.").
    output = json.loads(json.dumps(_CANNED_LLM_OUTPUT))
    output["parties"] = [
        {"role": "SELLER", "legalName": "ACME Corp.", "legalNameConfidence": 0.9, "taxIdentifier": None, "taxIdentifierConfidence": 0.0, "page": 1},
        {"role": "SELLER", "legalName": "ACME  Corp.", "legalNameConfidence": 0.99, "taxIdentifier": None, "taxIdentifierConfidence": 0.0, "page": 3},
        {"role": "BUYER", "legalName": "Beta Trading Ltd", "legalNameConfidence": 0.95, "taxIdentifier": None, "taxIdentifierConfidence": 0.0, "page": 1},
    ]
    _mock_llm(monkeypatch, output)

    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)
    validate_outgoing(event)

    parties = event["payload"]["result"]["parties"]
    assert len(parties) == 2  # iki ACME kaydi tek partiye birlesti
    seller = next(p for p in parties if p["role"] == "SELLER")
    assert seller["legalName"]["confidence"] == 0.99  # yuksek confidence'li versiyon tutuldu
    assert {r["page"] for r in seller["sourceReferences"]} == {1, 3}  # her iki sayfa da korundu

    warnings = event["payload"]["warnings"]
    assert any(w["code"] == "DUPLICATE_PARTY_MERGED" for w in warnings)


def test_money_rule_mapped_correctly(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _mock_llm(monkeypatch)
    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)
    rule = event["payload"]["result"]["rules"][0]
    assert rule["structuredValue"] == {"type": "MONEY", "amountMinor": 100000, "currency": "EUR"}


def test_structured_value_fallback_produces_warning(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    broken_output = json.loads(json.dumps(_CANNED_LLM_OUTPUT))
    broken_output["rules"][0]["amountMinor"] = None  # MONEY beyan edildi ama tutar yok
    broken_output["rules"][0]["currency"] = None
    _mock_llm(monkeypatch, broken_output)

    request = _request(base_url, "/pdf", _PDF_BYTES)
    event = run(request)  # yine schema-valid olmali (validate_outgoing pipeline icinde calisir)

    rule = event["payload"]["result"]["rules"][0]
    assert rule["structuredValue"]["type"] == "TEXT"
    warnings = event["payload"]["warnings"]
    assert any(w["code"] == "STRUCTURED_VALUE_FALLBACK" for w in warnings)


# --- OCR fallback (ADR-002 §3.1 "gerektiginde OCR") ---
# NOT: gercek Tesseract calisir - free/local/no-network oldugu icin GPT gibi
# mock-first kuralina tabi degil (ADR-004 §17'nin mock-first gerekcesi maliyet/
# network bagimliligidir; Tesseract'ta bunlarin hicbiri yok).

def test_scanned_pdf_triggers_real_ocr(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _mock_llm(monkeypatch)
    request = _request(base_url, "/scanned", _SCANNED_PDF_BYTES)
    event = run(request)

    validate_outgoing(event)
    document = event["payload"]["result"]["document"]
    assert document["textExtractionMethod"] == "OCR"

    warnings = event["payload"]["warnings"]
    assert any(w["code"] == "OCR_USED" for w in warnings)


def test_hybrid_pdf_mixes_digital_and_ocr(base_url: str, monkeypatch: pytest.MonkeyPatch) -> None:
    _mock_llm(monkeypatch)
    request = _request(base_url, "/hybrid", _HYBRID_PDF_BYTES)
    event = run(request)

    validate_outgoing(event)
    document = event["payload"]["result"]["document"]
    assert document["textExtractionMethod"] == "HYBRID"
    assert document["pageCount"] == 2


def test_blank_scanned_pdf_is_corrupted_file(base_url: str) -> None:
    # OCR sonrasi da hicbir metin cikmiyorsa (bombos taranmis sayfa) reddedilir.
    request = _request(base_url, "/blank-scanned", _BLANK_SCANNED_PDF_BYTES)
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.CORRUPTED_FILE


def test_ocr_extracts_real_readable_text(tmp_path: Path) -> None:
    from app.pipeline.document_extraction.text_extraction import extract_text

    pdf_path = tmp_path / "scanned.pdf"
    pdf_path.write_bytes(_SCANNED_PDF_BYTES)
    extracted = extract_text(pdf_path, detected_media_type=PDF)

    assert extracted.method == "OCR"
    assert extracted.ocr_pages == (1,)
    combined = extracted.full_text().upper()
    assert "PAYMENT" in combined
    assert "1000 EUR" in combined or "1000EUR" in combined.replace(" ", "")


# --- Hata yollari ---

def test_past_deadline_is_rejected_before_download(base_url: str) -> None:
    """Berke review #5: suresi gecmis deadline pahali indirme/LLM cagrisindan
    ONCE reddedilmeli. `/does-not-exist` gercekten fetch edilseydi farkli bir
    (deadline-disi) hata koduyla basarisiz olurdu -- INVALID_DEADLINE almamiz
    checkpoint'in indirmeden ONCE tetiklendigini kanitlar.
    """
    request = _request(base_url, "/does-not-exist", b"unused")
    request["payload"]["deadlineAt"] = "2000-01-01T00:00:00Z"
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.INVALID_DEADLINE


def test_hash_mismatch_fails_before_extraction(base_url: str) -> None:
    request = _request(base_url, "/pdf", _PDF_BYTES)
    request["payload"]["input"]["sha256"] = "a" * 64
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.CONTENT_HASH_MISMATCH


def test_unsupported_content_rejected(base_url: str) -> None:
    request = _request(base_url, "/junk", _JUNK)
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.UNSUPPORTED_MEDIA_TYPE


def test_declared_media_type_mismatch_rejected(base_url: str) -> None:
    request = _request(base_url, "/pdf", _PDF_BYTES, media_type=DOCX)
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.UNSUPPORTED_MEDIA_TYPE


def test_malformed_pdf_content_is_corrupted_file(base_url: str) -> None:
    # Dogru magic byte (%PDF-) ama gecersiz ic yapi: UNSUPPORTED_MEDIA_TYPE degil, CORRUPTED_FILE.
    request = _request(base_url, "/malformed-pdf", _MALFORMED_PDF)
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.CORRUPTED_FILE


def test_encrypted_pdf_is_rejected(base_url: str) -> None:
    request = _request(base_url, "/encrypted", _ENCRYPTED_PDF_BYTES)
    with pytest.raises(PipelineFailure) as exc:
        run(request)
    assert exc.value.code is ErrorCode.ENCRYPTED_DOCUMENT_UNSUPPORTED


# --- Giden event dogrulamasi (ADR-002 §11) ---

def test_unknown_outgoing_event_is_rejected() -> None:
    with pytest.raises(ContractViolation):
        validate_outgoing({"eventType": "ai.job.completed.v1", "jobType": "AUDIO_ANALYSIS"})


# --- Metin cikarimi birim testleri ---

def test_detect_media_type_reads_magic_bytes(tmp_path: Path) -> None:
    pdf = tmp_path / "a.pdf"
    pdf.write_bytes(_PDF_BYTES)
    assert detect_media_type(pdf, declared=PDF) == PDF

    docx_path = tmp_path / "a.docx"
    docx_path.write_bytes(_DOCX_BYTES)
    assert detect_media_type(docx_path, declared=DOCX) == DOCX


# --- Registry ---

def test_registry_routes_document_extraction_to_real_pipeline() -> None:
    assert pipeline_for("DOCUMENT_EXTRACTION") is run
