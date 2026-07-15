"""PDF/DOCX metin cikarimi + OCR fallback (ADR-002 SS3.1 "PDF veya DOCX metin
cikarimi, gerektiginde OCR").

Her PDF sayfasi once dijital metin katmanindan okunur. Sayfada yeterli metin
yoksa (taranmis/fotograflanmis sayfa oldugu varsayilir), sayfa goruntuye
cevrilip Tesseract ile OCR yapilir. Karisik belgeler (bazi sayfa dijital,
bazi sayfa taranmis) `textExtractionMethod: HYBRID` olarak isaretlenir
(ADR contract'inda zaten tanimli enum degeri).

Sayfa sinirlari korunur ki source-reference'lardaki `page` alani (ADR common
source-reference schema) gercek sayfaya isaret etsin.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader

from app.contracts.errors import ErrorCode, PipelineFailure
from app.pipeline.document_extraction.media import DOCX, PDF

# Sozlesmeler TR/EN karisik olabilir (ADR-002 ornegi: languageHints ["tr","en"]).
_OCR_LANGUAGES = "tur+eng"
# Bu esigin altindaki dijital metin "muhtemelen taranmis sayfa" kabul edilir.
_MIN_DIGITAL_CHARS = 20
# ~144 DPI render (2x zoom, 72 DPI taban) - OCR dogrulugu icin yeterli, hizli.
_OCR_RENDER_ZOOM = 2.0


@dataclass(frozen=True)
class ExtractedPage:
    page: int  # one-based (ADR source-reference schema)
    text: str


@dataclass(frozen=True)
class ExtractedDocument:
    pages: list[ExtractedPage]
    method: str  # canonical textExtractionMethod enum degeri
    ocr_pages: tuple[int, ...] = field(default_factory=tuple)

    @property
    def page_count(self) -> int:
        return len(self.pages)

    def full_text(self) -> str:
        """Her sayfayi acik [PAGE n] isaretiyle birlestirir (LLM'in dogru
        page numarasi raporlayabilmesi icin)."""
        return "\n\n".join(f"[PAGE {p.page}]\n{p.text}" for p in self.pages)


def _needs_ocr(digital_text: str) -> bool:
    return len((digital_text or "").strip()) < _MIN_DIGITAL_CHARS


def _ocr_page(doc: fitz.Document, page_index: int) -> str:
    try:
        page = doc.load_page(page_index)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(_OCR_RENDER_ZOOM, _OCR_RENDER_ZOOM))
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        return pytesseract.image_to_string(image, lang=_OCR_LANGUAGES)
    except Exception as exc:  # noqa: BLE001 - render veya OCR motoru hatasi
        raise PipelineFailure(
            ErrorCode.MODEL_PROVIDER_UNAVAILABLE,
            "OCR engine failed to process a page",
            details={"dependency": "tesseract", "reason": "ocr processing error"},
        ) from exc


def _extract_pdf(path: Path) -> ExtractedDocument:
    try:
        reader = PdfReader(str(path))
    except Exception as exc:  # noqa: BLE001 - kutuphane cesitli parse hatasi firlatabilir
        raise PipelineFailure(
            ErrorCode.CORRUPTED_FILE,
            "PDF content could not be parsed",
            details={"reason": "parse error"},
        ) from exc

    if reader.is_encrypted:
        raise PipelineFailure(
            ErrorCode.ENCRYPTED_DOCUMENT_UNSUPPORTED,
            "encrypted PDF is not supported",
            details={"reason": "encrypted"},
        )

    digital_texts = [(page.extract_text() or "") for page in reader.pages]
    if not digital_texts:
        raise PipelineFailure(
            ErrorCode.CORRUPTED_FILE,
            "PDF contains no pages",
            details={"reason": "empty document"},
        )

    pages: list[ExtractedPage] = []
    ocr_pages: list[int] = []
    fitz_doc: fitz.Document | None = None

    try:
        for i, digital_text in enumerate(digital_texts):
            page_number = i + 1
            if _needs_ocr(digital_text):
                if fitz_doc is None:
                    fitz_doc = fitz.open(str(path))
                ocr_text = _ocr_page(fitz_doc, i)
                if ocr_text.strip():
                    ocr_pages.append(page_number)
                pages.append(ExtractedPage(page=page_number, text=ocr_text))
            else:
                pages.append(ExtractedPage(page=page_number, text=digital_text))
    finally:
        if fitz_doc is not None:
            fitz_doc.close()

    if not any(p.text.strip() for p in pages):
        raise PipelineFailure(
            ErrorCode.CORRUPTED_FILE,
            "PDF contains no extractable text even after OCR",
            details={"reason": "empty after ocr"},
        )

    if not ocr_pages:
        method = "DIGITAL_PDF"
    elif len(ocr_pages) == len(pages):
        method = "OCR"
    else:
        method = "HYBRID"

    return ExtractedDocument(pages=pages, method=method, ocr_pages=tuple(ocr_pages))


def _extract_docx(path: Path) -> ExtractedDocument:
    try:
        document = DocxDocument(str(path))
    except Exception as exc:  # noqa: BLE001
        raise PipelineFailure(
            ErrorCode.CORRUPTED_FILE,
            "DOCX content could not be parsed",
            details={"reason": "parse error"},
        ) from exc

    text = "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip())
    if not text.strip():
        raise PipelineFailure(
            ErrorCode.CORRUPTED_FILE,
            "DOCX contains no extractable text",
            details={"reason": "empty document"},
        )
    # DOCX'te native sayfa siniri yok; tek mantiksal sayfa olarak ele alinir.
    # DOCX taranmis olamaz (inherently dijital format) - OCR fallback gerekmez.
    return ExtractedDocument(pages=[ExtractedPage(page=1, text=text)], method="DOCX")


def extract_text(path: Path, *, detected_media_type: str) -> ExtractedDocument:
    """Tespit edilen medya turune gore metni cikarir (gerekirse OCR ile).

    Firlatir: PipelineFailure(CORRUPTED_FILE / ENCRYPTED_DOCUMENT_UNSUPPORTED /
    MODEL_PROVIDER_UNAVAILABLE)
    """
    if detected_media_type == PDF:
        return _extract_pdf(path)
    if detected_media_type == DOCX:
        return _extract_docx(path)
    raise PipelineFailure(
        ErrorCode.UNSUPPORTED_MEDIA_TYPE,
        "no text extractor registered for this media type",
        details={"field": "input.mediaType", "reason": "unsupported format"},
    )
